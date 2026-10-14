"""Build the editable MX Gemmini overview from the BAR slide template.

Requires python-pptx. All diagrams, labels, callouts, and table cells remain
native PowerPoint objects; the university seal is the template's master image.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt


NAVY = RGBColor(16, 53, 82)
INK = RGBColor(49, 52, 85)
GREY = RGBColor(91, 89, 87)
CREAM = RGBColor(253, 248, 241)
PALE_BLUE = RGBColor(225, 232, 236)
SAND = RGBColor(237, 229, 220)
SAGE = RGBColor(229, 231, 218)
CODE_PURPLE = RGBColor(108, 61, 128)
CODE_GREEN = RGBColor(40, 112, 79)
CODE_ORANGE = RGBColor(164, 79, 37)
CODE_TEAL = RGBColor(0, 104, 116)

CODE_TOKEN = re.compile(
    r'"(?:[^"\\]|\\.)*"|%[A-Za-z][\w]*|\b(?:true|false|before|after|'
    r'__asm__|volatile|objdump|contract|return)\b|\b(?:0x[0-9a-fA-F]+|'
    r'\d+(?:\.\d+)?)\b|\b(?:mx_gemmini\.[A-Za-z_]+|'
    r'activation_format|weight_format|pe_mode|profile_sha256|'
    r'geometry|legal_compute|resources|mesh_rows|mesh_columns|'
    r'scratchpad_bytes|vpu|spad_requant)\b'
)


def code_color(token: str, following: str):
    if token.startswith('"'):
        if token.startswith('"mx_gemmini.'):
            return CODE_PURPLE
        return CODE_PURPLE if following.lstrip().startswith(":") else CODE_GREEN
    if token.startswith("%"):
        return CODE_TEAL
    if token[0].isdigit():
        return CODE_ORANGE
    return CODE_PURPLE


def set_text(shape, text: str, *, size: float, bold: bool = False,
             color=INK, align=PP_ALIGN.CENTER) -> None:
    frame = shape.text_frame
    frame.clear()
    frame.word_wrap = True
    frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    for index, line in enumerate(text.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = align
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        run = paragraph.add_run()
        run.text = line
        run.font.name = "Inter"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color


def add_text(slide, text: str, x: float, y: float, w: float, h: float,
             *, size: float, bold: bool = False, color=INK,
             align=PP_ALIGN.LEFT):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    shape.text_frame.margin_left = Inches(0.02)
    shape.text_frame.margin_right = Inches(0.02)
    shape.text_frame.margin_top = Inches(0.01)
    shape.text_frame.margin_bottom = Inches(0.01)
    set_text(shape, text, size=size, bold=bold, color=color, align=align)
    return shape


def add_panel(slide, x: float, y: float, w: float, h: float, fill):
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                   Inches(x), Inches(y), Inches(w), Inches(h))
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.color.rgb = RGBColor(176, 165, 151)
    return shape


def add_code(slide, source: str, x: float, y: float, w: float, h: float,
             *, size: float = 11.2):
    box = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                 Inches(x), Inches(y), Inches(w), Inches(h))
    box.fill.solid()
    box.fill.fore_color.rgb = CREAM
    box.line.color.rgb = RGBColor(211, 202, 191)
    frame = box.text_frame
    frame.clear()
    frame.word_wrap = False
    frame.vertical_anchor = MSO_ANCHOR.TOP
    frame.margin_left = Inches(0.12)
    frame.margin_top = Inches(0.10)
    frame.margin_right = Inches(0.08)
    frame.margin_bottom = Inches(0.04)
    for index, line in enumerate(source.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.alignment = PP_ALIGN.LEFT
        paragraph.space_before = Pt(0)
        paragraph.space_after = Pt(0)
        offset = 0
        for match in CODE_TOKEN.finditer(line):
            for value, color in ((line[offset:match.start()], NAVY),
                                 (match.group(), code_color(match.group(),
                                                            line[match.end():]))):
                if value:
                    run = paragraph.add_run()
                    run.text = value
                    run.font.name = "Liberation Mono"
                    run.font.size = Pt(size)
                    run.font.color.rgb = color
            offset = match.end()
        if offset < len(line):
            run = paragraph.add_run()
            run.text = line[offset:]
            run.font.name = "Liberation Mono"
            run.font.size = Pt(size)
            run.font.color.rgb = NAVY
    return box


def clear_content(slide) -> None:
    # Keep the template title. All new content remains native PowerPoint art.
    for shape in list(slide.shapes)[1:]:
        shape._element.getparent().remove(shape._element)


def keep_slides(prs: Presentation) -> None:
    # Template slides: process 20, architecture 23, results table 33.
    slide_ids = list(prs.slides._sldIdLst)
    retained = [slide_ids[index] for index in (22, 19, 32)]
    for slide_id in slide_ids:
        if slide_id not in retained:
            prs.part.drop_rel(slide_id.rId)
        prs.slides._sldIdLst.remove(slide_id)
    for slide_id in retained:
        prs.slides._sldIdLst.append(slide_id)


def style_master(prs: Presentation) -> None:
    master = prs.slides[0].slide_layout.slide_master
    by_name = {shape.name: shape for shape in master.shapes}
    set_text(by_name["FRAME | Lab name"], "BAR LAB", size=14, bold=True,
             color=NAVY, align=PP_ALIGN.LEFT)
    set_text(by_name["FRAME | Purpose label"], "PROGRAM", size=10,
             color=INK, align=PP_ALIGN.RIGHT)
    set_text(by_name["FRAME | Purpose value"], "MX GEMMINI + VPU", size=10,
             color=INK, align=PP_ALIGN.RIGHT)
    by_name["PAGE"].text = ""


def page_number(slide, number: int) -> None:
    add_text(slide, f"{number:02d} / 03", 11.45, 6.88, 1.35, 0.28,
             size=10, color=GREY, align=PP_ALIGN.RIGHT)


def architecture(slide) -> None:
    set_text(slide.shapes[0], "MX dialect: typed IR to RV64", size=27,
             bold=True)
    clear_content(slide)
    add_panel(slide, 0.94, 1.64, 5.65, 2.98, SAND)
    add_panel(slide, 6.76, 1.64, 5.65, 2.98, PALE_BLUE)
    add_text(slide, "FP4 GEMM → typed MX", 1.15, 1.82, 5.2, 0.34,
             size=17.5, bold=True, color=NAVY)
    add_code(slide,
             '%acc = "mx_gemmini.contract"(%A,%As,%B,%Bs) {\n'
             '  activation_format = "fp4_e2m1",\n'
             '  weight_format = "fp4_e2m1", pe_mode = 0 : i32 }\n'
             '%out = "mx_gemmini.readout_bf16"(%acc)',
             1.15, 2.24, 5.20, 1.42, size=11.1)
    add_text(slide,
             "resource → packed codes + E8M0 scales → mvin; contract → MX "
             "compute. Radiance softmax uses vpu_execute; readout_bf16 / "
             "spad_requant / wait finish the tile.",
             1.15, 3.76, 5.17, 0.75, size=10.6, color=INK)
    add_text(slide, "Compiled attention object → binary", 6.98, 1.82,
             5.2, 0.34, size=17.5, bold=True, color=NAVY)
    add_code(slide,
             '__asm__ volatile(\n'
             '  ".insn r 0x7b, 3, 2, x0, %0, %1"\n'
             '  : : "r"(Q_IN),\n'
             '      "r"(0x0010001000000000));\n'
             'objdump: 4a: 04b6b07b  .insn 4, 0x04b6b07b',
             6.98, 2.24, 5.20, 1.42, size=10.3)
    add_text(slide,
             "physical_program orders mvin → compute → VPU → mvout. "
             "RoCC C + data asm → RV64 GCC → ELF → Nicolas Spike.",
             6.98, 3.82, 5.17, 0.60, size=11.0, color=INK)

    add_text(slide, "Implemented transforms and checks", 0.96, 4.76,
             11.4, 0.30, size=16, bold=True, color=NAVY)
    transforms = (
        ("bind_profile", "Legal PE tuple +\nprofile digest"),
        ("MXBF16Catalog\nPreparation", "Select BF16\nMX sites"),
        ("schedule.plan", "Choose native\nloop policy"),
        ("MXVPUPostOffload", "Map softmax, add\nand reductions"),
        ("physical_program", "Order DMA, compute\nand VPU commands"),
    )
    for index, (name, action) in enumerate(transforms):
        x = 0.94 + index * 2.34
        add_panel(slide, x, 5.15, 2.15, 1.08,
                  (CREAM, PALE_BLUE, CREAM, SAGE, CREAM)[index])
        add_text(slide, name, x + 0.10, 5.25, 1.95, 0.37,
                 size=11.7, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
        add_text(slide, action, x + 0.10, 5.72, 1.95, 0.40,
                 size=10.8, color=INK, align=PP_ALIGN.CENTER)
    add_text(slide,
             "mx-gemmini-opt verifies typed IR. The standalone command path bypasses LLVM; "
             "the generic full-model host currently uses Merlin → LLVM → RV64.",
             0.96, 6.38, 11.45, 0.36, size=10.2, color=GREY)
    page_number(slide, 1)


def configurations(slide) -> None:
    set_text(slide.shapes[0], "One JSON profile changes legal MLIR", size=27,
             bold=True)
    clear_content(slide)
    add_panel(slide, 0.94, 1.62, 6.04, 4.82, SAND)
    add_panel(slide, 7.16, 1.62, 5.25, 4.82, PALE_BLUE)
    add_text(slide, "Selected RTL profile fields", 1.15, 1.79,
             5.58, 0.34, size=17.2, bold=True, color=NAVY)
    add_code(slide,
             '{\n'
             ' "geometry": {"mesh_rows":16,"mesh_columns":16},\n'
             ' "legal_compute": [\n'
             '  {"activation_format":"fp4_e2m1",\n'
             '   "weight_format":"fp4_e2m1","pe_mode":0},\n'
             '  {"activation_format":"fp8_e4m3",\n'
             '   "weight_format":"fp8_e4m3","pe_mode":8}\n'
             ' ],\n'
             ' "resources": {"scratchpad_bytes":262144,\n'
             '               "vpu":true,"spad_requant":true}\n'
             '}',
             1.15, 2.22, 5.57, 2.81, size=10.45)
    add_text(slide,
             "MxE4M3Fp4VpuGemminiRocketConfig · direct projections in both "
             "legal rows · two 8-lane BF16 VPUs.",
             1.15, 5.18, 5.57, 0.52, size=11.1, color=INK)
    add_text(slide,
             "RTL 266c593 + MxGen dba3e7 → 41 fragments / 40 SoC wrappers; "
             "each JSON stores source hashes.",
             1.15, 5.81, 5.57, 0.39, size=10.4, color=GREY)

    add_text(slide, "Before → after bind_profile", 7.38, 1.79,
             4.78, 0.34, size=17.2, bold=True, color=NAVY)
    add_code(slide,
             'before: contract {format = "mxfp4"}\n'
             'after:  contract {\n'
             '  activation_format = "fp4_e2m1",\n'
             '  weight_format = "fp4_e2m1",\n'
             '  pe_mode = 0 : i32,\n'
             '  profile_sha256 = "e9e33f..." }',
             7.38, 2.22, 4.79, 1.64, size=10.85)
    add_text(slide, "Compiler consequences", 7.38, 4.06,
             4.78, 0.30, size=14.1, bold=True, color=NAVY)
    add_text(slide,
             "• TinyLlama: 155 FP4 / 45 FP8 MX sites per stage.\n"
             "• FP4 binds PE 0; 256 KiB SPAD caps tiles.\n"
             "• VPU/requant legal; FP6 rejects in this profile.\n"
             "• load_profile(profile, rtl_root=rtl) rejects drift.",
             7.38, 4.42, 4.78, 1.62, size=11.6, color=INK)
    add_text(slide,
             "npu-exploration supplies source kernels for per-mode Spike checks; "
             "the RTL profile alone is structural, not a performance claim.",
             0.96, 6.50, 11.45, 0.25, size=10.2, color=GREY)
    page_number(slide, 2)


def set_cell(cell, text: str, *, heading: bool = False, bold: bool = False,
             align=PP_ALIGN.CENTER, size: float = 14.0) -> None:
    cell.text = ""
    cell.margin_left = Inches(0.14)
    cell.margin_right = Inches(0.10)
    cell.margin_top = Inches(0.05)
    cell.margin_bottom = Inches(0.03)
    cell.vertical_anchor = MSO_ANCHOR.MIDDLE
    paragraph = cell.text_frame.paragraphs[0]
    paragraph.alignment = align
    paragraph.space_before = Pt(0)
    paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    run.font.name = "Inter"
    run.font.size = Pt(size)
    run.font.bold = heading or bold
    run.font.color.rgb = RGBColor(255, 255, 255) if heading else INK


def results(slide, smol_comparison: dict | None,
            tiny_run: dict | None, tiny_comparison: dict | None) -> None:
    shapes = slide.shapes
    set_text(shapes[0], "Spike cycles by comparison scope", size=27,
             bold=True)
    frame = shapes[1]
    table = frame.table
    table.rows[0].height = Inches(0.60)
    for row in list(table.rows)[1:]:
        row.height = Inches(0.63)
    smol_compiler = (smol_comparison["reference_cycles"]
                     if smol_comparison else None)
    smol_native = (smol_comparison["candidate_cycles"]
                   if smol_comparison else None)
    tiny_compiler = (tiny_run["total_performance"]["cycles"]
                     if tiny_run else None)
    tiny_native = (tiny_comparison["candidate_cycles"]
                   if tiny_comparison else None)

    def delta(compiler: int | None, native: int | None) -> str:
        return f"{(compiler / native - 1) * 100:+.3f}%" if compiler and native else "—"

    values = (
        ("Workload / compared code", "Compiler cycles", "Nicolas path cycles", "Compiler Δ"),
        ("SmolVLA full 12 steps · issuer swap",
         f"{smol_compiler:,}*" if smol_compiler else "pending",
         f"{smol_native:,}" if smol_native else "running",
         delta(smol_compiler, smol_native)),
        ("TinyLlama full 33 steps · issuer swap",
         f"{tiny_compiler:,}*" if tiny_compiler else "running",
         f"{tiny_native:,}" if tiny_native else "pending",
         delta(tiny_compiler, tiny_native)),
        ("SmolVLA flow 1 step · issuer swap", "2,112,653,481", "2,112,919,925", "−0.0126%"),
        ("Dense attention 64×256×128 · fused source", "43,523", "43,641", "−0.27%"),
    )
    for row_index, row in enumerate(values):
        for col_index, text in enumerate(row):
            set_cell(table.cell(row_index, col_index), text,
                     heading=row_index == 0,
                     bold=row_index in (1, 2) and col_index == 0,
                     align=PP_ALIGN.LEFT if col_index == 0 else PP_ALIGN.CENTER,
                     size=11.2 if row_index == 0 else 12.4)
    add_shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.94),
                                       Inches(5.10), Inches(11.46), Inches(0.61))
    add_shape.fill.solid()
    add_shape.fill.fore_color.rgb = PALE_BLUE
    add_shape.line.fill.background()
    set_text(add_shape,
             "Full-model rows change only the MX matrix command code. Both runs use "
             "the same Merlin host and VPU code.\n"
             "Nicolas's fused MX+VPU kernels have not yet been stitched and timed "
             "as either complete model.",
             size=10.6, bold=True, color=NAVY, align=PP_ALIGN.CENTER)
    footer = shapes[2]
    footer.top = Inches(5.91)
    footer.height = Inches(0.66)
    if tiny_comparison and smol_comparison:
        footer_text = (
            "* Output hashes match: SmolVLA 33/33; TinyLlama 99/99. This checks "
            "the captured outputs, not model accuracy.\n"
            "Nicolas e4m3_vpu Spike, host tracking off. The dense attention "
            "row compares one fused kernel, not a whole model."
        )
    else:
        tiny_status = ("TinyLlama full receipt not supplied; "
                       "no model cycle total.")
        smol_status = ("Full SmolVLA issuer-swap outputs match (33/33)."
                       if smol_comparison else
                       "Full SmolVLA issuer-swap parity awaits its source run.")
        footer_text = (
            "* Completed timing receipt; independent quantized accuracy remains open. "
            f"{tiny_status}\n"
            "Nicolas e4m3_vpu Spike, host tracking off. Paired flow/attention "
            f"output hashes match. {smol_status}"
        )
    set_text(
        footer,
        footer_text,
        size=9.6, color=GREY, align=PP_ALIGN.LEFT,
    )
    page_number(slide, 3)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--template", type=Path, required=True)
    parser.add_argument("--out", type=Path, required=True)
    parser.add_argument("--smol-full-comparison", type=Path)
    parser.add_argument("--tiny-full-run", type=Path)
    parser.add_argument("--tiny-full-comparison", type=Path)
    args = parser.parse_args()
    if (args.tiny_full_run is None) != (args.tiny_full_comparison is None):
        parser.error("TinyLlama needs both the full run and equal-output comparison")
    smol_comparison = (json.loads(args.smol_full_comparison.read_text())
                       if args.smol_full_comparison else None)
    tiny_run = (json.loads(args.tiny_full_run.read_text())
                if args.tiny_full_run else None)
    tiny_comparison = (json.loads(args.tiny_full_comparison.read_text())
                       if args.tiny_full_comparison else None)
    if tiny_run is not None and (
            tiny_run.get("complete") is not True or tiny_run.get("exit_code") != 0
            or len(tiny_run.get("step_performance", [])) != 33
            or len(tiny_run.get("output_digests", [])) != 99
            or not isinstance((tiny_run.get("total_performance") or {}).get("cycles"), int)):
        parser.error("TinyLlama full run is incomplete")
    for comparison, expected in ((smol_comparison, None),
                                 (tiny_comparison,
                                  tiny_run["total_performance"]["cycles"]
                                  if tiny_run else None)):
        if comparison is None:
            continue
        if (comparison.get("schema") != "mx_gemmini.equal_input_session_cycles.v1"
                or comparison.get("output_digest_parity") is not True
                or not isinstance(comparison.get("reference_cycles"), int)
                or comparison["reference_cycles"] <= 0
                or (expected is not None and comparison["reference_cycles"] != expected)
                or not isinstance(comparison.get("candidate_cycles"), int)
                or comparison["candidate_cycles"] <= 0):
            parser.error("full-session comparison lacks equal-output timing evidence")
    if tiny_comparison is not None and (
            tiny_comparison.get("reference_run_sha256") !=
            hashlib.sha256(args.tiny_full_run.read_bytes()).hexdigest()):
        parser.error("TinyLlama comparison is not bound to the selected run")
    prs = Presentation(args.template)
    keep_slides(prs)
    style_master(prs)
    architecture(prs.slides[0])
    configurations(prs.slides[1])
    results(prs.slides[2], smol_comparison, tiny_run, tiny_comparison)
    args.out.parent.mkdir(parents=True, exist_ok=True)
    prs.save(args.out)


if __name__ == "__main__":
    main()
